#!/usr/bin/env python3
"""Find internships from the SimplifyJobs list and queue them in jobs.txt.

Source: the data file behind https://github.com/SimplifyJobs/Summer2027-Internships (updated daily).

Usage:
  python find_internships.py --folder "JOB_SEARCH_FOLDER" [filters] [--append]

Filters (all optional; without them, sensible defaults come from profile.json in the folder):
  --term "Summer 2027"        internship term (default: profile's internship term, else Summer 2027)
  --category software|data|hardware|product|quant   (repeatable; default: all)
  --location CA --location Seattle --location Remote  (repeatable; state codes, cities, SF, NYC, Remote)
  --keyword data --keyword analyst                  (repeatable; any must appear in the title)
  --exclude-keyword phd --exclude-keyword senior    (repeatable)
  --degree "Bachelor's"       keep listings open to this degree level (default: from profile.json)
  --needs-sponsorship         drop listings marked "Does Not Offer Sponsorship" or "U.S. Citizenship is Required".
                              Default from profile.json: if the user needs sponsorship, citizenship-only listings are
                              dropped; "Does Not Offer Sponsorship" listings are dropped only for seniors, since
                              underclassmen apply for the experience.
  --no-citizenship            drop "U.S. Citizenship is Required" listings
  --days 30                   only listings posted in the last N days
  --limit 25                  at most N results, newest first
  --append                    add the results to jobs.txt (above any "# done" line), one link per line with a note

Skips LinkedIn links (LinkedIn prohibits automated tools) and listings that are closed, already in tracker.xlsx (same company and role) or already in jobs.txt.
Prints one JSON object. Exit codes: 0 ok, 1 bad input, 6 couldn't download the list (network).
"""
import argparse
import datetime as dt
import json
import os
import re
import sys
import tempfile
import urllib.request

SOURCE = 'https://raw.githubusercontent.com/SimplifyJobs/Summer2027-Internships/dev/.github/scripts/listings.json'
REPO = 'https://github.com/SimplifyJobs/Summer2027-Internships'
CATEGORIES = {
    'software': {'Software', 'Software Engineering'},
    'data': {'AI/ML/Data', 'Data Science, AI & Machine Learning'},
    'hardware': {'Hardware', 'Hardware Engineering'},
    'product': {'Product', 'Product Management'},
    'quant': {'Quant', 'Quantitative Finance'},
}
for alias, target in (('ai', 'data'), ('ml', 'data'), ('swe', 'software'), ('pm', 'product')):
    CATEGORIES[alias] = CATEGORIES[target]
CITY_ALIASES = {
    'san francisco': ['sf', 'san francisco'], 'sf': ['sf', 'san francisco'],
    'new york': ['nyc', 'new york'], 'nyc': ['nyc', 'new york'],
    'bay area': ['sf', 'san francisco', 'san jose', 'santa clara', 'mountain view', 'palo alto', 'sunnyvale',
                 'menlo park', 'redwood city', 'cupertino', 'oakland', 'fremont', 'milpitas', 'san mateo', 'foster city'],
    'la': ['los angeles', 'la,'], 'los angeles': ['los angeles'],
}


def out(obj, code=0):
    print(json.dumps(obj, indent=1, default=str))
    sys.exit(code)


def norm(s):
    return re.sub(r'[^a-z0-9]+', ' ', str(s or '').lower()).strip()


def degree_level(text):
    t = norm(text)
    if re.search(r'\bph ?d\b|doctor', t):
        return 'PhD'
    if re.search(r'\bmba\b', t):
        return 'MBA'
    if re.search(r'master|\bm ?s\b|\bm ?sc\b|\bm ?eng\b|\bm ?a\b', t):
        return "Master's"
    if re.search(r'associate', t):
        return "Associate's"
    if re.search(r'bachelor|\bb ?s\b|\bb ?a\b|\bb ?sc\b|\bb ?eng\b|undergrad', t):
        return "Bachelor's"
    return None


def load_profile(folder):
    p = os.path.join(folder, 'profile.json')
    try:
        return json.load(open(p, encoding='utf-8'))
    except (OSError, ValueError):
        return {}


def below_senior(profile, term):
    """True for freshmen, sophomores and juniors; False for seniors and graduate students; None if unknown."""
    st = str(profile.get('classStanding') or '').lower()
    if st:
        if re.search(r'senior|fourth|4th|final|grad', st):
            return False
        if re.search(r'fresh|first|1st|soph|second|2nd|junior|third|3rd', st):
            return True
    grad = str(profile.get('expectedGraduation') or ((profile.get('education') or [{}])[0] or {}).get('end', ''))
    gy, ty = re.search(r'(20\d{2})', grad), re.search(r'(20\d{2})', term)
    if gy and ty:
        return int(gy.group(1)) > int(ty.group(1))
    return None


def location_match(locations, prefs):
    if not prefs:
        return True
    locs = [l.lower() for l in locations]
    for p in prefs:
        q = p.strip()
        if re.fullmatch(r'[A-Za-z]{2}', q) and q.upper() == q:  # state code like CA
            if any(re.search(rf',\s*{q.lower()}\b', l) for l in locs):
                return True
            continue
        ql = q.lower()
        if ql == 'remote':
            if any('remote' in l for l in locs):
                return True
            continue
        for alias in CITY_ALIASES.get(ql, [ql]):
            if any(alias == l or alias in l for l in locs):
                return True
    return False


def tracker_keys(folder):
    path = os.path.join(folder, 'tracker.xlsx')
    keys = set()
    if not os.path.exists(path):
        return keys
    try:
        from openpyxl import load_workbook
        ws = load_workbook(path, read_only=True).active
        for i, row in enumerate(ws.iter_rows(values_only=True)):
            if i == 0 or not row:
                continue
            keys.add((norm(row[2]), norm(row[3])))  # company, role
    except Exception:
        pass
    return keys


def jobs_urls(path):
    if not os.path.exists(path):
        return set(), []
    lines = open(path, encoding='utf-8').read().splitlines()
    urls = {m.group(0).split('?')[0].rstrip('/') for l in lines for m in [re.search(r'https?://\S+', l)] if m}
    return urls, lines


def fetch():
    cache = os.path.join(tempfile.gettempdir(), 'simplify-summer2027-listings.json')
    try:
        if os.path.exists(cache) and (dt.datetime.now().timestamp() - os.path.getmtime(cache)) < 3 * 3600:
            return json.load(open(cache, encoding='utf-8'))
        req = urllib.request.Request(SOURCE, headers={'User-Agent': 'internship-application-skill'})
        with urllib.request.urlopen(req, timeout=60) as r:
            data = r.read()
        open(cache, 'wb').write(data)  # temp folder only
        return json.loads(data)
    except Exception as e:
        out({'ok': False, 'code': 'NETWORK', 'error': f'Could not download the SimplifyJobs list: {e}',
             'hint': f'Allow raw.githubusercontent.com in network settings, or open {REPO} in the browser and read the table there.'}, 6)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--folder', default='.')
    ap.add_argument('--term')
    ap.add_argument('--category', action='append', default=[])
    ap.add_argument('--location', action='append', default=[])
    ap.add_argument('--keyword', action='append', default=[])
    ap.add_argument('--exclude-keyword', action='append', default=[])
    ap.add_argument('--degree')
    ap.add_argument('--needs-sponsorship', action='store_true', default=None)
    ap.add_argument('--no-citizenship', action='store_true')
    ap.add_argument('--days', type=int)
    ap.add_argument('--limit', type=int, default=25)
    ap.add_argument('--append', action='store_true')
    a = ap.parse_args()

    folder = os.path.abspath(a.folder)
    profile = load_profile(folder)
    defaults_used = {}

    term = a.term
    if not term:
        m = re.search(r'(summer|fall|winter|spring)\s+(20\d{2})', str(profile.get('internshipTerm', '')), re.I)
        term = f'{m.group(1).title()} {m.group(2)}' if m else 'Summer 2027'
        defaults_used['term'] = term

    degree = a.degree
    if not degree:
        edu = (profile.get('education') or [{}])[0]
        degree = degree_level(edu.get('degree', ''))
        if degree:
            defaults_used['degree'] = degree
    forced = a.needs_sponsorship is True
    needs_sponsorship = forced or profile.get('needsSponsorship') == 'yes'
    below = below_senior(profile, term)
    # Before senior year the user applies for experience, so no-sponsorship postings stay in.
    drop_no_sponsor = needs_sponsorship and (forced or below is False)
    drop_citizenship = needs_sponsorship or a.no_citizenship  # needing sponsorship means not a citizen
    if not forced and needs_sponsorship:
        defaults_used['sponsorship'] = ('keeping postings that do not sponsor (below senior)' if not drop_no_sponsor
                                        else 'dropping postings that do not sponsor (senior)')

    cats = set()
    for c in a.category:
        key = c.strip().lower()
        if key not in CATEGORIES:
            out({'ok': False, 'error': f'unknown category: {c}', 'categories': ['software', 'data', 'hardware', 'product', 'quant']}, 1)
        cats |= CATEGORIES[key]

    data = fetch()
    cutoff = (dt.datetime.now() - dt.timedelta(days=a.days)).timestamp() if a.days else None
    tracked = tracker_keys(folder)
    jobs_path = os.path.join(folder, 'jobs.txt')
    queued, lines = jobs_urls(jobs_path)
    kw = [norm(k) for k in a.keyword]
    ex = [norm(k) for k in a.exclude_keyword]

    skipped = {'linkedin_link': 0, 'closed': 0, 'term': 0, 'category': 0, 'location': 0, 'keyword': 0, 'degree': 0,
               'sponsorship': 0, 'too_old': 0, 'already_tracked': 0, 'already_in_jobs_txt': 0}
    results = []
    for x in data:
        if 'linkedin.com' in (x.get('url') or '').lower():  # never queue LinkedIn; apply on the company's site
            skipped['linkedin_link'] += 1; continue
        if not x.get('active') or not x.get('is_visible', True):
            skipped['closed'] += 1; continue
        if term not in (x.get('terms') or []):
            skipped['term'] += 1; continue
        if cats and x.get('category') not in cats:
            skipped['category'] += 1; continue
        if not location_match(x.get('locations') or [], a.location):
            skipped['location'] += 1; continue
        title = norm(x.get('title'))
        if (kw and not any(k in title for k in kw)) or any(k in title for k in ex):
            skipped['keyword'] += 1; continue
        degrees = x.get('degrees') or []
        if degree and degrees and degree not in degrees:
            skipped['degree'] += 1; continue
        sp = x.get('sponsorship', '')
        if (drop_no_sponsor and sp == 'Does Not Offer Sponsorship') or \
                (drop_citizenship and sp == 'U.S. Citizenship is Required'):
            skipped['sponsorship'] += 1; continue
        if cutoff and (x.get('date_posted') or 0) < cutoff:
            skipped['too_old'] += 1; continue
        if (norm(x.get('company_name')), title) in tracked:
            skipped['already_tracked'] += 1; continue
        if x.get('url', '').split('?')[0].rstrip('/') in queued:
            skipped['already_in_jobs_txt'] += 1; continue
        results.append(x)

    results.sort(key=lambda x: -(x.get('date_posted') or 0))
    seen, unique = set(), []
    for x in results:  # the same role is sometimes listed twice under different links
        key = (norm(x.get('company_name')), norm(x.get('title')))
        if key not in seen:
            seen.add(key)
            unique.append(x)
    skipped['duplicate_listing'] = len(results) - len(unique)
    results = unique
    total_matches = len(results)
    results = results[:a.limit]
    listings = [{
        'company': x['company_name'], 'role': x['title'], 'locations': x.get('locations', []),
        'posted': dt.date.fromtimestamp(x['date_posted']).isoformat() if x.get('date_posted') else '',
        'category': x.get('category'), 'degrees': x.get('degrees', []), 'sponsorship': x.get('sponsorship'),
        'url': x.get('url'),
    } for x in results]

    added = 0
    if a.append and listings:
        new_lines = [f"{l['url']}  # {l['company']} | {l['role']} | {', '.join(l['locations'][:3])} | posted {l['posted']}"
                     for l in listings]
        done_at = next((i for i, l in enumerate(lines) if l.strip().lower() == '# done'), len(lines))
        pending, rest = lines[:done_at], lines[done_at:]
        while pending and not pending[-1].strip():
            pending.pop()
        merged = pending + new_lines + ([''] if rest else []) + rest
        with open(jobs_path, 'w', encoding='utf-8') as f:  # rewrites jobs.txt in place; nothing is deleted
            f.write('\n'.join(merged) + '\n')
        added = len(new_lines)

    out({'ok': True, 'source': REPO, 'term': term, 'filters': {
        'categories': a.category, 'locations': a.location, 'keywords': a.keyword, 'exclude_keywords': a.exclude_keyword,
        'degree': degree, 'drop_no_sponsorship_postings': drop_no_sponsor,
        'drop_citizenship_required': drop_citizenship, 'below_senior': below, 'days': a.days}, 'defaults_from_profile': defaults_used,
        'matches': total_matches, 'shown': len(listings), 'added_to_jobs_txt': added, 'skipped': skipped, 'listings': listings})


if __name__ == '__main__':
    main()
