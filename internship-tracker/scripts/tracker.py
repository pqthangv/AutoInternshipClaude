#!/usr/bin/env python3
"""Keep the internship tracker spreadsheet (tracker.xlsx) in the job-search folder.

Usage:
  python tracker.py add     FILE --set company=Acme --set role="Data Intern" [--set key=value ...]
  python tracker.py update  FILE --id A0003 --set status=Submitted [--set key=value ...]
  python tracker.py find    FILE [--id A0003] [--company acme] [--role data] [--status Submitted]
  python tracker.py summary FILE

Keys: id, date_added, company, role, term, location, deadline, status, status_updated,
      next_step, interview_date, posting_url, folder, resume, gaps, notes

- add refuses duplicates (same company, role and term) and returns the existing row instead.
- update sets status_updated to today whenever status changes.
- The file is created with headers, a status dropdown, colors and filters if it doesn't exist.
  An older tracker.csv beside it is imported once and then moved into _old/ (never deleted).
- Nothing is ever deleted. Prints one JSON object.

Exit codes: 0 ok, 1 bad input or not found, 4 the file is open in another program (close Excel and retry),
            5 openpyxl is missing (pip install openpyxl).
"""
import argparse
import csv
import datetime as dt
import json
import os
import re
import shutil
import sys

COLUMNS = [
    ('id', 'ID', 8), ('date_added', 'Date added', 12), ('company', 'Company', 22), ('role', 'Role', 34),
    ('term', 'Term', 13), ('location', 'Location', 20), ('deadline', 'Deadline', 12), ('status', 'Status', 17),
    ('status_updated', 'Status updated', 14), ('next_step', 'Next step', 28), ('interview_date', 'Interview date', 15),
    ('posting_url', 'Posting', 12), ('folder', 'Folder', 12), ('resume', 'Resume file', 30),
    ('gaps', 'Gaps', 30), ('notes', 'Notes', 40),
]
KEYS = [c[0] for c in COLUMNS]
STATUSES = ['Ready to submit', 'Submitted', 'Online assessment', 'Interviewing', 'Offer', 'Rejected',
            'Withdrawn', 'Needs you', 'Skipped']
COLORS = {  # soft fills that read in Excel's light and dark modes
    'Ready to submit': 'FFF4CC', 'Submitted': 'DCE9FB', 'Online assessment': 'E8DDF7', 'Interviewing': 'D5F0E1',
    'Offer': 'B9E6C9', 'Rejected': 'EDEDED', 'Withdrawn': 'EDEDED', 'Needs you': 'FBE0CF', 'Skipped': 'F3F3F3',
}
CSV_ALIASES = {'date': 'date_added', 'resume_file': 'resume'}


def out(obj, code=0):
    print(json.dumps(obj, indent=1, default=str))
    sys.exit(code)


try:
    from openpyxl import Workbook, load_workbook
    from openpyxl.formatting.rule import FormulaRule
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation
except ImportError:
    out({'ok': False, 'code': 'NO_OPENPYXL', 'error': 'openpyxl is not installed. Run: pip install openpyxl'}, 5)


def norm(s):
    return re.sub(r'[^a-z0-9]+', ' ', str(s or '').lower()).strip()


def canonical_status(value):
    v = norm(value)
    for s in STATUSES:
        if norm(s) == v:
            return s
    for s in STATUSES:  # "needs you: login" -> Needs you, "skipped: not eligible" -> Skipped
        if v.startswith(norm(s)):
            return s
    if v in ('applied', 'sent'):
        return 'Submitted'
    if v in ('oa', 'assessment', 'coding challenge'):
        return 'Online assessment'
    if v in ('interview', 'interviews'):
        return 'Interviewing'
    return None


def create(path):
    wb = Workbook()
    ws = wb.active
    ws.title = 'Applications'
    for i, (_, header, width) in enumerate(COLUMNS, start=1):
        c = ws.cell(row=1, column=i, value=header)
        c.font = Font(bold=True, color='FFFFFF')
        c.fill = PatternFill('solid', fgColor='17213A')
        c.alignment = Alignment(vertical='center')
        ws.column_dimensions[c.column_letter].width = width
    ws.freeze_panes = 'C2'
    ws.auto_filter.ref = f'A1:{ws.cell(row=1, column=len(COLUMNS)).column_letter}1'
    col = ws.cell(row=1, column=KEYS.index('status') + 1).column_letter
    dv = DataValidation(type='list', formula1='"' + ','.join(STATUSES) + '"', allow_blank=True)
    dv.error = 'Pick a status from the list.'
    ws.add_data_validation(dv)
    dv.add(f'{col}2:{col}2000')
    for status, color in COLORS.items():
        ws.conditional_formatting.add(f'A2:{ws.cell(row=1, column=len(COLUMNS)).column_letter}2000',
                                      FormulaRule(formula=[f'${col}2="{status}"'], fill=PatternFill('solid', fgColor=color)))
    return wb


def rows(ws):
    data = []
    for r in range(2, ws.max_row + 1):
        vals = [ws.cell(row=r, column=i + 1).value for i in range(len(KEYS))]
        if any(v not in (None, '') for v in vals):
            row = {k: ('' if v is None else v) for k, v in zip(KEYS, vals)}
            row['_row'] = r
            data.append(row)
    return data


def write_row(ws, r, row, base_dir):
    for i, k in enumerate(KEYS, start=1):
        v = row.get(k, '')
        c = ws.cell(row=r, column=i, value=v if v != '' else None)
        c.alignment = Alignment(vertical='top', wrap_text=k in ('gaps', 'notes', 'next_step', 'role'))
        if k == 'posting_url' and v:
            c.hyperlink = str(v)
            c.value = 'Open posting'
            c.style = 'Hyperlink'
        elif k == 'folder' and v:
            c.hyperlink = str(v).replace('\\', '/')  # relative link opens the job folder from Excel
            c.value = 'Open folder'
            c.style = 'Hyperlink'


def read_links(ws, row):
    """Hyperlink cells show a label; the real value is the link target."""
    for k in ('posting_url', 'folder'):
        c = ws.cell(row=row['_row'], column=KEYS.index(k) + 1)
        if c.hyperlink is not None and c.hyperlink.target:
            row[k] = c.hyperlink.target
    return row


def next_id(existing):
    nums = [int(m.group(1)) for r in existing for m in [re.match(r'A(\d+)$', str(r.get('id', '')))] if m]
    return f'A{(max(nums) + 1) if nums else 1:04d}'


def move_to_old(path):
    folder = os.path.join(os.path.dirname(path), '_old')
    os.makedirs(folder, exist_ok=True)
    stamp = dt.datetime.now().strftime('%Y%m%d-%H%M%S')
    base, ext = os.path.splitext(os.path.basename(path))
    dest = os.path.join(folder, f'{base}-{stamp}{ext}')
    shutil.move(path, dest)
    return dest


def open_book(path):
    base_dir = os.path.dirname(os.path.abspath(path))
    notes = []
    if os.path.exists(path):
        wb = load_workbook(path)
        return wb, wb['Applications'] if 'Applications' in wb.sheetnames else wb.active, notes
    wb = create(path)
    ws = wb.active
    legacy = os.path.join(base_dir, 'tracker.csv')
    if os.path.exists(legacy):
        with open(legacy, newline='', encoding='utf-8-sig') as f:
            imported = []
            for rec in csv.DictReader(f):
                row = {CSV_ALIASES.get(k, k): (v or '').strip() for k, v in rec.items() if k}
                raw = row.get('status', '')
                row['status'] = canonical_status(raw) or raw
                if ':' in raw and not row.get('next_step'):
                    row['next_step'] = raw.split(':', 1)[1].strip()
                row['id'] = next_id(imported)
                imported.append(row)
        for i, row in enumerate(imported, start=2):
            write_row(ws, i, row, base_dir)
        notes.append(f'Imported {len(imported)} rows from tracker.csv and moved it to {move_to_old(legacy)}')
    return wb, ws, notes


def save(wb, path):
    try:
        wb.save(path)
    except PermissionError:
        out({'ok': False, 'code': 'FILE_IN_USE', 'error': f'{os.path.basename(path)} is open in another program. Close it in Excel and run again.'}, 4)


def parse_sets(pairs):
    data = {}
    for p in pairs or []:
        if '=' not in p:
            out({'ok': False, 'error': f'--set needs key=value, got: {p}'}, 1)
        k, v = p.split('=', 1)
        k = k.strip().lower()
        if k not in KEYS:
            out({'ok': False, 'error': f'unknown key: {k}', 'keys': KEYS}, 1)
        data[k] = v.strip()
    if 'status' in data and data['status']:
        s = canonical_status(data['status'])
        if not s:
            out({'ok': False, 'error': f'unknown status: {data["status"]}', 'statuses': STATUSES}, 1)
        # keep the reason ("needs you: login") in next_step when the status carried one
        reason = data['status'].split(':', 1)[1].strip() if ':' in data['status'] else ''
        data['status'] = s
        if reason:
            data['next_step'] = data.get('next_step') or reason
    return data


def clean(row):
    return {k: v for k, v in row.items() if not k.startswith('_')}


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('command', choices=['add', 'update', 'find', 'summary'])
    ap.add_argument('file')
    ap.add_argument('--set', action='append')
    ap.add_argument('--id')
    ap.add_argument('--company')
    ap.add_argument('--role')
    ap.add_argument('--status')
    a = ap.parse_args()

    path = os.path.abspath(a.file)
    base_dir = os.path.dirname(path)
    wb, ws, notes = open_book(path)
    existing = [read_links(ws, r) for r in rows(ws)]
    today = dt.date.today().isoformat()

    if a.command == 'add':
        data = parse_sets(a.set)
        if not data.get('company') or not data.get('role'):
            out({'ok': False, 'error': 'add needs at least company and role'}, 1)
        for r in existing:
            if norm(r['company']) == norm(data['company']) and norm(r['role']) == norm(data['role']) \
                    and norm(r.get('term')) == norm(data.get('term', r.get('term'))):
                if notes:
                    save(wb, path)
                out({'ok': True, 'created': False, 'duplicate': True, 'row': clean(r), 'notes': notes})
        data.setdefault('date_added', today)
        data.setdefault('status', 'Ready to submit')
        data.setdefault('status_updated', today)
        data['id'] = next_id(existing)
        r = (max([x['_row'] for x in existing]) + 1) if existing else 2
        write_row(ws, r, data, base_dir)
        save(wb, path)
        out({'ok': True, 'created': True, 'row': {k: data.get(k, '') for k in KEYS}, 'notes': notes})

    if a.command == 'update':
        if not a.id:
            out({'ok': False, 'error': 'update needs --id (use find to look it up)'}, 1)
        match = [r for r in existing if str(r['id']).upper() == a.id.upper()]
        if not match:
            out({'ok': False, 'error': f'no row with id {a.id}'}, 1)
        row = match[0]
        data = parse_sets(a.set)
        if 'status' in data and data['status'] != row.get('status'):
            data.setdefault('status_updated', today)
        row.update(data)
        write_row(ws, row['_row'], row, base_dir)
        save(wb, path)
        out({'ok': True, 'row': clean(row), 'notes': notes})

    if notes:
        save(wb, path)

    if a.command == 'find':
        res = existing
        if a.id:
            res = [r for r in res if str(r['id']).upper() == a.id.upper()]
        if a.company:
            res = [r for r in res if norm(a.company) in norm(r['company'])]
        if a.role:
            res = [r for r in res if norm(a.role) in norm(r['role'])]
        if a.status:
            s = canonical_status(a.status) or a.status
            res = [r for r in res if r.get('status') == s]
        out({'ok': True, 'count': len(res), 'rows': [clean(r) for r in res], 'notes': notes})

    if a.command == 'summary':
        counts = {s: 0 for s in STATUSES}
        for r in existing:
            if r.get('status') in counts:
                counts[r['status']] += 1
        def as_date(v):
            if isinstance(v, (dt.date, dt.datetime)):
                return v if isinstance(v, dt.date) and not isinstance(v, dt.datetime) else v.date()
            try:
                return dt.date.fromisoformat(str(v)[:10])
            except ValueError:
                return None
        upcoming = sorted([clean(r) for r in existing if as_date(r.get('interview_date')) and as_date(r['interview_date']) >= dt.date.today()],
                          key=lambda r: as_date(r['interview_date']))
        not_sent = [clean(r) for r in existing if r.get('status') == 'Ready to submit']
        waiting = [clean(r) for r in existing if r.get('status') == 'Submitted' and as_date(r.get('status_updated'))
                   and (dt.date.today() - as_date(r['status_updated'])).days >= 21]
        out({'ok': True, 'total': len(existing), 'by_status': {k: v for k, v in counts.items() if v},
             'upcoming_interviews': upcoming, 'filled_not_submitted': not_sent,
             'submitted_3_weeks_no_update': waiting, 'notes': notes})


if __name__ == '__main__':
    main()
